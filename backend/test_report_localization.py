from __future__ import annotations

import inspect
import re
import unittest
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from urllib.parse import unquote

from fontTools.ttLib import TTFont
from openpyxl import load_workbook
from pptx import Presentation

from app.routers import reports
from app.routers.config import CompanyUpdateRequest, MineUpdateRequest
from app.models.company import CompanySettings
from app.models.mine import MineSettings
from app.services import (
    daily_executive_pdf_service,
    excel_service,
    monthly_kpi_pdf_service,
    powerpoint_service,
    weekly_operations_pdf_service,
)
from app.services.daily_executive_pdf_service import (
    _daily_executive_brief,
    generate_daily_executive_pdf,
)
from app.services.report_localization import (
    build_content_disposition,
    build_report_filename,
    ensure_report_pdf_fonts,
    format_report_date,
    get_report_display_name,
    localize_report_label,
    normalize_report_language,
    pluralize_report_count,
)
from app.services.monthly_kpi_pdf_service import (
    _chart_value as _monthly_chart_value,
    _event_phrase as _monthly_event_phrase,
    _monthly_mn_executive_brief,
    _monthly_safety_requires_attention,
    _monthly_safety_status,
    _monthly_mn_safety_action,
    _monthly_mn_safety_outlook,
    generate_monthly_kpi_pdf,
)
from app.services.report_branding_service import (
    ReportBranding,
    resolve_customer_display_identity,
    resolve_report_branding,
)
from app.services.weekly_operations_pdf_service import (
    _event_phrase as _weekly_event_phrase,
    _weekly_mn_executive_brief,
    _trend_sentence as _weekly_trend_sentence,
    generate_weekly_operations_pdf,
)


class ReportLocalizationTests(unittest.TestCase):
    def test_language_normalization(self):
        expected = {
            "EN": "en",
            "en": "en",
            "MN": "mn",
            "mn": "mn",
            "mongolian": "mn",
            "монгол": "mn",
            "invalid": "en",
            None: "en",
        }
        for value, language in expected.items():
            with self.subTest(value=value):
                self.assertEqual(normalize_report_language(value), language)

    def test_shared_semantic_labels(self):
        self.assertEqual(localize_report_label("stable", "en"), "Stable")
        self.assertEqual(localize_report_label("stable", "mn"), "Тогтвортой")
        self.assertEqual(localize_report_label("watch", "mn"), "Анхаарах")
        self.assertEqual(localize_report_label("critical", "mn"), "Ноцтой")
        self.assertEqual(localize_report_label("declining", "EN"), "Declining")
        self.assertEqual(localize_report_label("declining", "MN"), "Буурч байна")
        self.assertEqual(localize_report_label("no_data", "mn"), "Өгөгдөлгүй")
        self.assertEqual(localize_report_label("above_target", "mn"), "Зорилтоос дээгүүр")
        semantic_labels = {
            "high": "Өндөр",
            "medium": "Дунд",
            "low": "Бага",
            "open": "Нээлттэй",
            "in_progress": "Хэрэгжиж Байна",
            "not_started": "Эхлээгүй",
            "completed": "Дууссан",
            "overdue": "Хугацаа Хэтэрсэн",
            "unassigned": "Хариуцагчгүй",
            "not_set": "Тодорхойлоогүй",
        }
        for key, mongolian in semantic_labels.items():
            with self.subTest(key=key):
                self.assertEqual(localize_report_label(key, "mn"), mongolian)

    def test_daily_common_labels_in_both_languages(self):
        expected = {
            "daily_report": ("Daily Executive Report", "Өдрийн Удирдлагын Тайлан"),
            "mine_health": ("Mine Health", "Уурхайн Үйл Ажиллагааны Үнэлгээ"),
            "today_at_a_glance": ("Today at a Glance", "Өнөөдрийн Товч Үзүүлэлт"),
            "management_attention": (
                "Management Attention",
                "Удирдлагын Анхаарах Асуудал",
            ),
            "executive_brief": ("Executive Brief", "Удирдлагын Товч Дүгнэлт"),
            "todays_actions": ("Today's Actions", "Өнөөдрийн Арга Хэмжээ"),
            "operation_status": ("Operation Status", "Үйл Ажиллагааны Төлөв"),
            "production": ("Production", "Үйлдвэрлэл"),
            "plant": ("Plant", "Боловсруулах Үйлдвэр"),
            "safety": ("Safety", "Аюулгүй Ажиллагаа"),
            "impact": ("Impact", "Нөлөөлөл"),
            "action": ("Action", "Арга Хэмжээ"),
        }
        for key, (english, mongolian) in expected.items():
            with self.subTest(key=key):
                self.assertEqual(localize_report_label(key, "en"), english)
                self.assertEqual(localize_report_label(key, "mn"), mongolian)

    def test_customer_identity_resolver_is_configured_and_backward_compatible(self):
        scenarios = (
            (
                {
                    "company_name": "Achit-Ikht LLC",
                    "company_name_en": "Achit-Ikht LLC",
                    "company_name_mn": None,
                    "mine_name": "Achit-Ikht Copper Cathode Operation",
                    "mine_name_en": "Achit-Ikht Copper Cathode Operation",
                    "mine_name_mn": "Ачит-Ихт Зэсийн Катодын Үйлдвэр",
                },
                ("Achit-Ikht LLC", "Achit-Ikht Copper Cathode Operation"),
                ("Achit-Ikht LLC", "Ачит-Ихт Зэсийн Катодын Үйлдвэр"),
            ),
            (
                {
                    "company_name": "Example Copper Ltd",
                    "company_name_en": "Example Copper Ltd",
                    "company_name_mn": "Жишээ Зэс ХХК",
                    "mine_name": "Example Copper Operation",
                    "mine_name_en": "Example Copper Operation",
                    "mine_name_mn": "Жишээ Зэсийн Үйл Ажиллагаа",
                },
                ("Example Copper Ltd", "Example Copper Operation"),
                ("Жишээ Зэс ХХК", "Жишээ Зэсийн Үйл Ажиллагаа"),
            ),
            (
                {
                    "company_name": "North Mining",
                    "mine_name": "North Mine Operation",
                    "company_name_en": None,
                    "company_name_mn": None,
                    "mine_name_en": None,
                    "mine_name_mn": None,
                },
                ("North Mining", "North Mine Operation"),
                ("North Mining", "North Mine Operation"),
            ),
        )

        for configuration, expected_en, expected_mn in scenarios:
            with self.subTest(configuration=configuration):
                english = resolve_customer_display_identity(configuration, "en")
                mongolian = resolve_customer_display_identity(configuration, "mn")
                omitted = resolve_customer_display_identity(configuration)
                invalid = resolve_customer_display_identity(configuration, "invalid")
                self.assertEqual(
                    (english.company_name, english.operation_name), expected_en
                )
                self.assertEqual(
                    (mongolian.company_name, mongolian.operation_name), expected_mn
                )
                self.assertEqual(
                    (omitted.company_name, omitted.operation_name), expected_en
                )
                self.assertEqual(
                    (invalid.company_name, invalid.operation_name), expected_en
                )

    def test_customer_identity_is_not_hardcoded_in_generic_localization(self):
        localization_source = Path(
            "app/services/report_localization.py"
        ).read_text(encoding="utf-8")
        self.assertNotIn("Achit-Ikht", localization_source)
        self.assertNotIn("Ачит-Ихт", localization_source)

    def test_all_artifacts_use_the_shared_customer_identity_boundary(self):
        direct_resolver_generators = (
            generate_daily_executive_pdf,
            generate_weekly_operations_pdf,
            generate_monthly_kpi_pdf,
        )
        for generator in direct_resolver_generators:
            self.assertIn(
                "resolve_customer_display_identity",
                inspect.getsource(generator),
            )
        for generator in (
            excel_service.generate_executive_excel_export,
            powerpoint_service.generate_executive_powerpoint,
        ):
            self.assertIn("resolve_report_branding", inspect.getsource(generator))

    def test_bilingual_identity_fields_are_exposed_by_models_and_config_api(self):
        self.assertTrue(
            {"company_name_en", "company_name_mn"}.issubset(
                CompanySettings.__table__.columns.keys()
            )
        )
        self.assertTrue(
            {"mine_name_en", "mine_name_mn"}.issubset(
                MineSettings.__table__.columns.keys()
            )
        )
        company_fields = CompanyUpdateRequest.model_fields
        mine_fields = MineUpdateRequest.model_fields
        self.assertTrue(
            {"company_name_en", "company_name_mn"}.issubset(company_fields)
        )
        self.assertTrue(
            {"mine_name_en", "mine_name_mn"}.issubset(mine_fields)
        )

    def test_report_branding_resolution_preserves_source_configuration(self):
        source = ReportBranding(
            company_name="Example Copper Ltd",
            mine_name="Example Copper Operation",
            company_name_en="Example Copper Ltd",
            company_name_mn="Жишээ Зэс ХХК",
            mine_name_en="Example Copper Operation",
            mine_name_mn="Жишээ Зэсийн Үйл Ажиллагаа",
            logo_url=None,
            logo_path=None,
            primary_color="#0F172A",
            secondary_color="#020617",
            timezone="Asia/Ulaanbaatar",
            language="English",
        )
        mongolian = resolve_report_branding(source, "mn")
        self.assertEqual(mongolian.company_name, "Жишээ Зэс ХХК")
        self.assertEqual(
            mongolian.mine_name,
            "Жишээ Зэсийн Үйл Ажиллагаа",
        )
        self.assertEqual(source.company_name, "Example Copper Ltd")
        self.assertEqual(source.mine_name, "Example Copper Operation")

    def test_report_names_and_pluralization(self):
        self.assertEqual(
            get_report_display_name("daily_pdf", "en"),
            "Daily Executive Report",
        )
        self.assertEqual(
            get_report_display_name("daily_pdf", "mn"),
            "Өдрийн удирдлагын тайлан",
        )
        self.assertEqual(
            pluralize_report_count(1, singular="action"),
            "1 action",
        )
        self.assertEqual(
            pluralize_report_count(2, singular="action"),
            "2 actions",
        )
        self.assertEqual(
            pluralize_report_count(
                2,
                language="mn",
                singular="action",
                mongolian_label="арга хэмжээ",
            ),
            "2 арга хэмжээ",
        )

    def test_all_localized_filenames(self):
        generated_on = datetime(2026, 9, 8, 14, 5)
        expected = {
            "daily_pdf": (
                "Daily_Executive_Report_2026-09-08.pdf",
                "Өдрийн_Удирдлагын_Тайлан_2026-09-08.pdf",
            ),
            "weekly_pdf": (
                "Weekly_Operations_Report_2026-09-08.pdf",
                "Долоо_Хоногийн_Үйл_Ажиллагааны_Тайлан_2026-09-08.pdf",
            ),
            "monthly_pdf": (
                "Monthly_KPI_Pack_2026-09-08.pdf",
                "Сарын_KPI_Тайлан_2026-09-08.pdf",
            ),
            "excel_export": (
                "Mine_Manager_AI_Executive_Export_2026-09-08.xlsx",
                "Уурхайн_Менежер_AI_Удирдлагын_Экспорт_2026-09-08.xlsx",
            ),
            "board_pack": (
                "Mine_Manager_AI_Executive_Board_Pack_2026-09-08.pptx",
                "Уурхайн_Менежер_AI_Удирдлагын_Танилцуулга_2026-09-08.pptx",
            ),
        }
        for report_type, (english, mongolian) in expected.items():
            with self.subTest(report_type=report_type):
                self.assertEqual(build_report_filename(report_type, "en", generated_on), english)
                self.assertEqual(build_report_filename(report_type, "mn", generated_on), mongolian)

    def test_content_disposition_preserves_unicode_and_sanitizes_paths(self):
        filename = "ӨҮөү/../Тайлан\\2026.pdf"
        fallback = "Daily/Executive\\Report.pdf"
        header = build_content_disposition(filename, fallback)

        self.assertIn('filename="Daily_Executive_Report.pdf"', header)
        self.assertIn("filename*=UTF-8''", header)
        encoded = header.split("filename*=UTF-8''", 1)[1]
        decoded = unquote(encoded)
        self.assertIn("ӨҮөү", decoded)
        self.assertNotIn("/", decoded)
        self.assertNotIn("\\", decoded)

    def test_date_formats_are_deterministic(self):
        value = datetime(2026, 9, 8, 14, 5)
        self.assertEqual(format_report_date(value, "en"), "08 Sep 2026")
        self.assertEqual(
            format_report_date(value, "en", "generated_timestamp"),
            "08 Sep 2026 14:05",
        )
        self.assertEqual(format_report_date(value, "mn"), "2026 оны 9-р сарын 8")
        self.assertEqual(
            format_report_date(value, "mn", "generated_timestamp"),
            "2026 оны 9-р сарын 8 14:05",
        )
        self.assertEqual(
            format_report_date(value, "mn", "short_chart_date"),
            "9-р сарын 8",
        )
        self.assertEqual(value, datetime(2026, 9, 8, 14, 5))

    def test_mongolian_pdf_fonts_resolve_and_cover_required_glyphs(self):
        fonts = ensure_report_pdf_fonts("mn")
        for path_key in ("regular_path", "bold_path"):
            path = Path(fonts[path_key])
            self.assertTrue(path.is_file())
            cmap = TTFont(path).getBestCmap()
            for character in "АЯаяӨөҮү":
                self.assertIn(ord(character), cmap)

    def test_daily_deterministic_narratives(self):
        english = _daily_executive_brief(
            ["Production", "Plant"],
            safety_status="Stable",
            incidents=0,
            near_misses=1,
            critical_risks=0,
            language="en",
        )
        mongolian = _daily_executive_brief(
            ["Үйлдвэрлэл", "Боловсруулах Үйлдвэр"],
            safety_status="Stable",
            incidents=0,
            near_misses=1,
            critical_risks=0,
            language="mn",
        )
        self.assertEqual(
            english,
            "Production and Plant require management attention today. "
            "Safety remains stable with no reported incidents or critical risks.",
        )
        self.assertEqual(
            mongolian,
            "Өнөөдөр үйлдвэрлэл болон боловсруулах үйлдвэрийн гүйцэтгэлд "
            "удирдлагын анхаарал шаардлагатай байна. Аюулгүй ажиллагаа тогтвортой "
            "бөгөөд осол болон ноцтой эрсдэл бүртгэгдээгүй.",
        )
        production_only = _daily_executive_brief(
            ["Үйлдвэрлэл"],
            safety_status="Stable",
            incidents=0,
            near_misses=0,
            critical_risks=0,
            language="mn",
        )
        self.assertTrue(production_only.startswith("Өнөөдөр үйлдвэрлэлд "))

    def test_daily_en_and_mn_smoke_generation_remains_one_page(self):
        daily_data = {
            "mine_name": "Achit-Ikht Copper Cathode Operation",
            "mine_name_en": "Achit-Ikht Copper Cathode Operation",
            "mine_name_mn": "Ачит-Ихт Зэсийн Катодын Үйлдвэр",
            "operation_profile": "sxew_copper",
            "fleet_applicable": False,
            "production_label": "Cathode Production",
            "report_date": "2026-09-08",
            "health": 93.5,
            "ore": 87.9,
            "plant": 94.6,
            "throughput": 97.3,
            "recovery": 90.5,
            "safety": 0,
            "safety_score": 98.0,
            "near_misses": 1,
            "critical_risks": 0,
        }
        original = dict(daily_data)
        with patch.object(
            daily_executive_pdf_service,
            "_report_header",
            wraps=daily_executive_pdf_service._report_header,
        ) as header:
            buffers = (
                generate_daily_executive_pdf(daily_data),
                generate_daily_executive_pdf(daily_data, language="en"),
                generate_daily_executive_pdf(daily_data, language="mn"),
            )
        self.assertEqual(
            header.call_args.kwargs["mine_name"],
            "Ачит-Ихт Зэсийн Катодын Үйлдвэр",
        )
        page_counts = [
            len(re.findall(rb"/Type\s*/Page(?!s)", buffer.getvalue()))
            for buffer in buffers
        ]
        self.assertEqual(page_counts, [1, 1, 1])
        self.assertEqual(daily_data, original)

    def test_weekly_static_labels_and_semantic_terms(self):
        expected = {
            "weekly_report": "Долоо Хоногийн Үйл Ажиллагааны Тайлан",
            "weekly_mine_health": "Долоо Хоногийн Уурхайн Үйл Ажиллагааны Үнэлгээ",
            "weekly_performance_at_a_glance": "Долоо Хоногийн Товч Гүйцэтгэл",
            "weekly_trends": "Долоо Хоногийн Хандлага",
            "next_week_actions": "Ирэх Долоо Хоногийн Арга Хэмжээ",
            "next_week_outlook": "Ирэх Долоо Хоногийн Төлөв",
            "weekly_kpi_detail": "Долоо Хоногийн KPI Дэлгэрэнгүй",
            "weekly_operating_status": "Долоо Хоногийн Үйл Ажиллагааны Төлөв",
            "plant_throughput": "Боловсруулах Үйлдвэрийн Нэвтрүүлэх Чадвар",
            "plant_recovery": "Металл Авалт",
            "result": "Үр Дүн",
        }
        for key, mongolian in expected.items():
            with self.subTest(key=key):
                self.assertEqual(localize_report_label(key, "mn"), mongolian)

    def test_weekly_trends_events_and_counts_are_localized(self):
        self.assertEqual(_weekly_trend_sentence("Declining"), "Declined during the reporting period")
        self.assertEqual(
            _weekly_trend_sentence("Declining", "mn"),
            "Тайлант хугацаанд гүйцэтгэл буурсан.",
        )
        self.assertEqual(_weekly_event_phrase(0, "incident", language="mn"), "осол бүртгэгдээгүй")
        self.assertEqual(_weekly_event_phrase(1, "near miss", language="mn"), "1 осолд дөхсөн тохиолдол")
        self.assertEqual(_weekly_event_phrase(2, "critical risk", language="mn"), "2 ноцтой эрсдэл")
        self.assertEqual(
            pluralize_report_count(3, language="mn", singular="item", mongolian_label="асуудал"),
            "3 асуудал",
        )

    def test_weekly_mn_executive_brief_is_dynamic_and_safety_conditional(self):
        cases = {
            ("production",): (
                "Үйлдвэрлэлд удирдлагын анхаарал шаардлагатай байна."
            ),
            ("production", "plant_throughput"): (
                "Үйлдвэрлэл болон боловсруулах үйлдвэрийн нэвтрүүлэх чадварт "
                "удирдлагын анхаарал шаардлагатай байна."
            ),
            ("safety",): (
                "Аюулгүй ажиллагаанд удирдлагын анхаарал шаардлагатай байна. "
                "Аюулгүй ажиллагааны хяналтыг ирэх долоо хоногт идэвхтэй "
                "үргэлжлүүлэх шаардлагатай."
            ),
            ("production", "plant_throughput", "safety"): (
                "Үйлдвэрлэл, боловсруулах үйлдвэрийн нэвтрүүлэх чадвар болон "
                "аюулгүй ажиллагаанд удирдлагын анхаарал шаардлагатай байна. "
                "Аюулгүй ажиллагааны хяналтыг ирэх долоо хоногт идэвхтэй "
                "үргэлжлүүлэх шаардлагатай."
            ),
        }
        for areas, expected in cases.items():
            with self.subTest(areas=areas):
                self.assertEqual(_weekly_mn_executive_brief(list(areas)), expected)

        without_safety = _weekly_mn_executive_brief(["plant_recovery"])
        self.assertNotIn("Аюулгүй ажиллагааны хяналтыг", without_safety)

    def test_weekly_en_and_mn_smoke_generation_remains_two_pages(self):
        days = [
            {
                "report_date": f"2026-09-0{day}",
                "ore": 94 - day,
                "throughput": 99 - day / 10,
                "recovery": 92 - day / 5,
            }
            for day in range(2, 9)
        ]
        weekly_data = {
            "mine_name": "Achit-Ikht Copper Cathode Operation",
            "mine_name_en": "Achit-Ikht Copper Cathode Operation",
            "mine_name_mn": "Ачит-Ихт Зэсийн Катодын Үйлдвэр",
            "operation_profile": "sxew_copper",
            "production_label": "Cathode Production",
            "period_start": "2026-09-02",
            "period_end": "2026-09-08",
            "health": 94.0,
            "ore": 87.4,
            "throughput": 97.6,
            "recovery": 90.8,
            "safety": 0,
            "safety_score": 99.7,
            "near_misses": 1,
            "critical_risks": 0,
            "days": days,
        }
        original = {**weekly_data, "days": [dict(day) for day in days]}
        with patch.object(
            weekly_operations_pdf_service,
            "_report_header",
            wraps=weekly_operations_pdf_service._report_header,
        ) as header:
            buffers = (
                generate_weekly_operations_pdf(weekly_data),
                generate_weekly_operations_pdf(weekly_data, language="en"),
                generate_weekly_operations_pdf(weekly_data, language="mn"),
            )
        self.assertEqual(
            header.call_args.kwargs["mine_name"],
            "Ачит-Ихт Зэсийн Катодын Үйлдвэр",
        )
        page_counts = [
            len(re.findall(rb"/Type\s*/Page(?!s)", buffer.getvalue()))
            for buffer in buffers
        ]
        self.assertEqual(page_counts, [2, 2, 2])
        self.assertEqual(weekly_data, original)

    def test_monthly_static_labels_and_semantic_terms(self):
        expected = {
            "monthly_report": "Сарын KPI Тайлан",
            "monthly_mine_health": "Уурхайн Үйл Ажиллагааны Үнэлгээ",
            "monthly_performance_at_a_glance": "Сарын Товч Гүйцэтгэл",
            "monthly_trends": "Сарын Гүйцэтгэлийн Хандлага",
            "strategic_actions": "Стратегийн Арга Хэмжээ",
            "next_month_outlook": "Ирэх Сарын Төлөв",
            "monthly_kpi_detail": "KPI Дэлгэрэнгүй",
            "monthly_operating_status": "Үйл Ажиллагааны Төлөв",
            "result": "Үр Дүн",
            "variance": "Зөрүү",
            "actual": "Бодит Гүйцэтгэл",
            "events": "Үйл Явдал",
        }
        for key, mongolian in expected.items():
            with self.subTest(key=key):
                self.assertEqual(localize_report_label(key, "mn"), mongolian)

    def test_monthly_mn_executive_brief_dynamic_combinations(self):
        combinations = (
            ["production"],
            ["plant_throughput"],
            ["plant_recovery"],
            ["safety"],
            ["production", "plant_throughput"],
            ["production", "plant_recovery"],
            ["production", "plant_throughput", "safety"],
            ["production", "plant_throughput", "plant_recovery", "safety"],
            [],
        )
        for areas in combinations:
            with self.subTest(areas=areas):
                brief = _monthly_mn_executive_brief(areas)
                self.assertTrue(brief.endswith("."))
                self.assertNotIn("Боловсруулах Үйлдвэрийн", brief)
                self.assertEqual(
                    "Аюулгүй ажиллагааны эрсдэлийн хяналтыг" in brief,
                    "safety" in areas,
                )

        self.assertEqual(
            _monthly_mn_executive_brief(["production"]),
            "Үйлдвэрлэлд удирдлагын анхаарал шаардлагатай байна.",
        )
        self.assertEqual(
            _monthly_mn_executive_brief(["production", "plant_throughput"]),
            "Үйлдвэрлэл болон боловсруулах үйлдвэрийн нэвтрүүлэх чадварт "
            "удирдлагын анхаарал шаардлагатай байна.",
        )

    def test_monthly_safety_events_and_missing_chart_values(self):
        self.assertEqual(_monthly_event_phrase(0, "incident", language="mn"), "0 осол")
        self.assertEqual(
            _monthly_event_phrase(1, "near miss", language="mn"),
            "1 осолд дөхсөн тохиолдол",
        )
        self.assertEqual(
            _monthly_event_phrase(1, "critical risk", language="mn"),
            "1 ноцтой эрсдэл",
        )
        self.assertIsNone(_monthly_chart_value({"throughput_chart": None}, "throughput_chart", "throughput"))
        self.assertEqual(_monthly_chart_value({"throughput": 97.2}, "throughput_chart", "throughput"), 97.2)

        self.assertEqual(_monthly_safety_status(0, 0, 99.6), "On Track")
        self.assertTrue(_monthly_safety_requires_attention(0, 1, 0, 99.6))
        self.assertEqual(_monthly_safety_status(0, 1, 99.6), "Watch")
        self.assertTrue(_monthly_safety_requires_attention(0, 0, 1, 99.6))
        self.assertFalse(_monthly_safety_requires_attention(0, 0, 0, 99.6))

    def test_monthly_mn_safety_actions_follow_actual_event_types(self):
        near_miss_only = _monthly_mn_safety_action(0, 3, 0)
        self.assertEqual(
            near_miss_only,
            "Аюулгүй ажиллагааны эрсдэлийн хяналтыг баталгаажуулж, осолд "
            "дөхсөн тохиолдлуудыг шалган, шаардлагатай залруулах арга хэмжээг "
            "хэрэгжүүлнэ.",
        )
        cases = {
            "critical_only": _monthly_mn_safety_action(0, 0, 1),
            "incident_only": _monthly_mn_safety_action(1, 0, 0),
            "near_and_critical": _monthly_mn_safety_action(0, 2, 1),
            "incident_and_near": _monthly_mn_safety_action(1, 2, 0),
        }
        expected_tokens = {
            "critical_only": ("ноцтой эрсдэлийг",),
            "incident_only": ("ослыг",),
            "near_and_critical": ("тохиолдлуудыг", "ноцтой эрсдэлийг"),
            "incident_and_near": ("ослыг", "тохиолдлуудыг"),
        }
        absent_tokens = {
            "critical_only": ("ослыг", "тохиолдлыг"),
            "incident_only": ("тохиолдлыг", "ноцтой эрсдэлийг"),
            "near_and_critical": ("ослыг",),
            "incident_and_near": ("ноцтой эрсдэлийг",),
        }
        for name, action in cases.items():
            with self.subTest(name=name):
                self.assertTrue(all(token in action for token in expected_tokens[name]))
                self.assertTrue(all(token not in action for token in absent_tokens[name]))

        self.assertIsNone(
            _monthly_mn_safety_action(0, 0, 0, safety_exception=False)
        )
        self.assertEqual(
            _monthly_mn_safety_outlook(True, 0),
            "Аюулгүй ажиллагааны залруулах арга хэмжээг хэрэгжүүлж, эрсдэлийн "
            "хяналтыг баталгаажуулах",
        )
        self.assertEqual(
            _monthly_mn_safety_outlook(False, 0),
            "Аюулгүй ажиллагааны хяналтыг тогтвортой үргэлжлүүлэх",
        )

    def test_monthly_en_and_mn_smoke_generation_remains_two_pages(self):
        days = []
        for day in range(1, 31):
            days.append({
                "report_date": f"2026-08-{day:02d}",
                "ore": 92 - day / 10,
                "throughput_chart": None if day in {5, 18} else 98 - day / 30,
                "recovery_chart": None if day in {7, 19} else 91 - day / 50,
            })
        monthly_data = {
            "mine_name": "Achit-Ikht Copper Cathode Operation",
            "mine_name_en": "Achit-Ikht Copper Cathode Operation",
            "mine_name_mn": "Ачит-Ихт Зэсийн Катодын Үйлдвэр",
            "operation_profile": "sxew_copper",
            "production_label": "Cathode Production",
            "period_start": "2026-08-01",
            "period_end": "2026-08-30",
            "health": 92.0,
            "ore": 87.9,
            "throughput": 97.6,
            "recovery": 88.8,
            "safety": 0,
            "safety_score": 99.6,
            "near_misses": 1,
            "critical_risks": 1,
            "days": days,
        }
        original = {**monthly_data, "days": [dict(day) for day in days]}
        with patch.object(
            monthly_kpi_pdf_service,
            "_report_header",
            wraps=monthly_kpi_pdf_service._report_header,
        ) as header:
            buffers = (
                generate_monthly_kpi_pdf(monthly_data),
                generate_monthly_kpi_pdf(monthly_data, language="en"),
                generate_monthly_kpi_pdf(monthly_data, language="mn"),
            )
        self.assertEqual(
            header.call_args.args[0],
            "Ачит-Ихт Зэсийн Катодын Үйлдвэр",
        )
        page_counts = [
            len(re.findall(rb"/Type\s*/Page(?!s)", buffer.getvalue()))
            for buffer in buffers
        ]
        self.assertEqual(page_counts, [2, 2, 2])
        self.assertEqual(monthly_data, original)


class ExcelLocalizationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.branding = ReportBranding(
            company_name="Achit-Ikht LLC",
            mine_name="Achit-Ikht Copper Cathode Operation",
            company_name_en="Achit-Ikht LLC",
            mine_name_en="Achit-Ikht Copper Cathode Operation",
            mine_name_mn="Ачит-Ихт Зэсийн Катодын Үйлдвэр",
            logo_url=None,
            logo_path=None,
            primary_color="#0F172A",
            secondary_color="#020617",
            timezone="Asia/Ulaanbaatar",
            language="English",
            company_id=2,
            mine_id=2,
        )
        created_at = datetime(2026, 9, 9, 8, 30)
        cls.production_rows = [
            {
                "report_date": date(2026, 9, day),
                "mine_name": "Achit-Ikht Copper Cathode Operation",
                "ore_plan": 50000 + day,
                "ore_actual": 45000 + day,
                "waste_plan": 0,
                "waste_actual": 0,
                "created_at": created_at,
            }
            for day in range(2, 9)
        ]
        cls.plant_rows = [
            {
                "report_date": date(2026, 9, day),
                "mine_name": "Achit-Ikht Copper Cathode Operation",
                "throughput_plan": 45000 + day,
                "throughput_actual": 44000 + day,
                "recovery": 90.0 + day / 10,
                "created_at": created_at,
            }
            for day in range(2, 9)
        ]
        cls.safety_rows = [
            {
                "report_date": date(2026, 9, day),
                "mine_name": "Achit-Ikht Copper Cathode Operation",
                "incidents": 0,
                "near_misses": 1 if day == 8 else 0,
                "critical_risks": 0,
                "safety_score": 98.0 if day == 8 else 100.0,
                "created_at": created_at,
            }
            for day in range(2, 9)
        ]

    def _generate(self, language_marker=inspect.Parameter.empty):
        with (
            patch.object(excel_service, "get_report_branding", return_value=self.branding),
            patch.object(excel_service, "_fetch_production_data", return_value=self.production_rows),
            patch.object(excel_service, "_fetch_plant_data", return_value=self.plant_rows),
            patch.object(excel_service, "_fetch_safety_data", return_value=self.safety_rows),
        ):
            kwargs = {
                "db": object(),
                "company_id": 2,
                "mine_id": 2,
                "operation_profile": "sxew_copper",
            }
            if language_marker is not inspect.Parameter.empty:
                kwargs["language"] = language_marker
            return load_workbook(excel_service.generate_executive_excel_export(**kwargs))

    @staticmethod
    def _sheet_signature(worksheet):
        return {
            "dimensions": (worksheet.max_row, worksheet.max_column),
            "merges": tuple(sorted(str(item) for item in worksheet.merged_cells.ranges)),
            "freeze": str(worksheet.freeze_panes),
            "filter": worksheet.auto_filter.ref,
            "formulas": tuple(
                (cell.coordinate, cell.value)
                for row in worksheet.iter_rows()
                for cell in row
                if cell.data_type == "f"
            ),
        }

    def test_excel_default_and_explicit_english_preserve_baseline_structure(self):
        default = self._generate()
        english = self._generate("en")
        expected_names = [
            "Executive Summary",
            "Cathode Production",
            "Plant",
            "Safety",
            "KPI Definitions",
        ]
        expected_dimensions = [(18, 6), (11, 7), (11, 8), (11, 7), (11, 5)]
        self.assertEqual(default.sheetnames, expected_names)
        self.assertEqual(english.sheetnames, expected_names)
        self.assertEqual(
            [(sheet.max_row, sheet.max_column) for sheet in english.worksheets],
            expected_dimensions,
        )
        for default_sheet, english_sheet in zip(default.worksheets, english.worksheets):
            self.assertEqual(self._sheet_signature(default_sheet), self._sheet_signature(english_sheet))
            self.assertEqual(
                [[cell.value for cell in row] for row in default_sheet.iter_rows()],
                [[cell.value for cell in row] for row in english_sheet.iter_rows()],
            )

    def test_excel_mn_sheet_names_labels_types_and_semantic_parity(self):
        english = self._generate("en")
        mongolian = self._generate("mn")
        expected_names = [
            "Удирдлагын Товч",
            "Катодын Үйлдвэрлэл",
            "Боловсруулах Үйлдвэр",
            "Аюулгүй Ажиллагаа",
            "KPI Тайлбар",
        ]
        self.assertEqual(mongolian.sheetnames, expected_names)
        self.assertEqual(len(set(expected_names)), 5)
        self.assertTrue(all(len(name) <= 31 for name in expected_names))
        self.assertTrue(all(not re.search(r"[:\\/?*\[\]]", name) for name in expected_names))

        summary, production, plant, safety, definitions = mongolian.worksheets
        self.assertEqual(summary["A1"].value, "Удирдлагын Үйл Ажиллагааны Товч")
        self.assertEqual(summary["A4"].value, "Компани")
        self.assertEqual(summary["A5"].value, "Үйл ажиллагаа")
        self.assertEqual(summary["B5"].value, "Ачит-Ихт Зэсийн Катодын Үйлдвэр")
        self.assertEqual(summary["A15"].value, "Өгөгдлийн Хамрах Хүрээ")
        self.assertEqual(production["A4"].value, "Огноо")
        self.assertEqual(production["C4"].value, "Үйлдвэрлэлийн Төлөвлөгөө")
        self.assertEqual(plant["C4"].value, "Нэвтрүүлэх Чадварын Төлөвлөгөө")
        self.assertEqual(plant["G4"].value, "Металл авалт")
        self.assertEqual(safety["C4"].value, "Осол")
        self.assertEqual(safety["D4"].value, "Осолд Дөхсөн Тохиолдол")
        self.assertEqual(safety["E4"].value, "Ноцтой Эрсдэл")
        self.assertEqual(definitions["C4"].value, "Тодорхойлолт")

        self.assertEqual(summary["C11"].value, "Анхаарах")
        self.assertEqual(summary["C12"].value, "Анхаарах")
        self.assertEqual(summary["C13"].value, "Анхаарах")

        for english_sheet, mongolian_sheet in zip(english.worksheets, mongolian.worksheets):
            self.assertEqual(self._sheet_signature(english_sheet), self._sheet_signature(mongolian_sheet))
            for english_row, mongolian_row in zip(
                english_sheet.iter_rows(),
                mongolian_sheet.iter_rows(),
            ):
                for english_cell, mongolian_cell in zip(english_row, mongolian_row):
                    if isinstance(english_cell.value, (int, float, date, datetime)):
                        self.assertEqual(english_cell.value, mongolian_cell.value)
                        self.assertIs(type(english_cell.value), type(mongolian_cell.value))
                        self.assertEqual(english_cell.number_format, mongolian_cell.number_format)

        self.assertIsInstance(production["A5"].value, datetime)
        self.assertIsInstance(production["C5"].value, (int, float))
        self.assertIsInstance(safety["C5"].value, int)
        self.assertEqual(production["A5"].number_format, "yyyy-mm-dd")
        self.assertEqual(production["G5"].number_format, "yyyy-mm-dd hh:mm")

        all_text = "\n".join(
            str(cell.value)
            for sheet in mongolian.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str)
        )
        forbidden = (
            "Executive Summary",
            "Cathode Production",
            "Plant Performance",
            "Safety Performance",
            "KPI Definitions",
            "Data Coverage",
            "Report Date",
            "Created At",
            "Near Misses",
            "Critical Risks",
            "Below target",
            "No data",
        )
        self.assertTrue(all(token not in all_text for token in forbidden))
        self.assertTrue(all(character in all_text for character in "ӨөҮү"))
        self.assertEqual(production["A1"].font.name, "Arial")


class PowerPointLocalizationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.branding = ReportBranding(
            company_name="Achit-Ikht LLC",
            mine_name="Achit-Ikht Copper Cathode Operation",
            company_name_en="Achit-Ikht LLC",
            mine_name_en="Achit-Ikht Copper Cathode Operation",
            mine_name_mn="Ачит-Ихт Зэсийн Катодын Үйлдвэр",
            logo_url=None,
            logo_path=None,
            primary_color="#0F172A",
            secondary_color="#020617",
            timezone="Asia/Ulaanbaatar",
            language="English",
            company_id=2,
            mine_id=2,
        )
        cls.production_rows = [
            {
                "report_date": date(2026, 8, 26) + timedelta(days=index),
                "ore_plan": 52000 + index * 100,
                "ore_actual": 46000 + index * 50,
            }
            for index in range(14)
        ]
        cls.plant_rows = [
            {
                "report_date": date(2026, 8, 30) + timedelta(days=index),
                "throughput_plan": 45000,
                "throughput_actual": 43800,
                "recovery": 90.5,
            }
            for index in range(10)
        ]
        cls.safety_rows = [
            {
                "report_date": date(2026, 9, 8),
                "incidents": 0,
                "near_misses": 1,
                "critical_risks": 0,
                "safety_score": 98.0,
            }
        ]
        cls.actions = [
            {
                "title": "Review the lowest-performing operational KPI",
                "priority": "high",
                "owner": "Үйл ажиллагаа",
                "timing": None,
                "status": "In_Progress",
            },
            {
                "title": "Investigate production loss versus plan",
                "priority": "medium",
                "owner": "Processing Operations",
                "timing": None,
                "status": "Completed",
            },
        ]

    def _generate(self, language_marker=inspect.Parameter.empty):
        with (
            patch.object(powerpoint_service, "get_report_branding", return_value=self.branding),
            patch.object(powerpoint_service, "_fetch_production_data", return_value=self.production_rows),
            patch.object(powerpoint_service, "_fetch_fleet_data", return_value=[]),
            patch.object(powerpoint_service, "_fetch_plant_data", return_value=self.plant_rows),
            patch.object(powerpoint_service, "_fetch_safety_data", return_value=self.safety_rows),
            patch.object(powerpoint_service, "_fetch_executive_actions", return_value=self.actions),
        ):
            kwargs = {
                "db": object(),
                "company_id": 2,
                "mine_id": 2,
                "operation_profile": "sxew_copper",
            }
            if language_marker is not inspect.Parameter.empty:
                kwargs["language"] = language_marker
            return Presentation(powerpoint_service.generate_executive_powerpoint(**kwargs))

    @staticmethod
    def _slide_text(slide):
        return "\n".join(
            shape.text
            for shape in slide.shapes
            if hasattr(shape, "text_frame") and shape.text
        )

    @staticmethod
    def _chart_signature(presentation):
        signature = []
        for slide_number, slide in enumerate(presentation.slides, 1):
            for shape in slide.shapes:
                if not getattr(shape, "has_chart", False):
                    continue
                chart = shape.chart
                signature.append(
                    (
                        slide_number,
                        len(chart.plots[0].categories),
                        tuple(tuple(series.values) for series in chart.series),
                    )
                )
        return tuple(signature)

    def test_powerpoint_default_and_explicit_english_preserve_structure(self):
        default = self._generate()
        english = self._generate("en")
        expected_titles = (
            "Executive Operations\nBoard Pack",
            "Executive KPI Summary",
            "Cathode Production Trend",
            "Plant and Safety Overview",
            "Key Operational Risks",
            "Management Action Register",
            "Executive Recommendations",
        )
        self.assertEqual(len(default.slides), 7)
        self.assertEqual(len(english.slides), 7)
        for index, title in enumerate(expected_titles):
            self.assertIn(title, self._slide_text(default.slides[index]))
            self.assertIn(title, self._slide_text(english.slides[index]))
        self.assertEqual(self._chart_signature(default), self._chart_signature(english))

    def test_powerpoint_mn_titles_statuses_and_cyrillic(self):
        presentation = self._generate("mn")
        expected_titles = (
            "Удирдлагын Үйл Ажиллагааны Танилцуулга",
            "Удирдлагын KPI Товч",
            "Катодын Үйлдвэрлэлийн Хандлага",
            "Боловсруулах Үйлдвэр ба Аюулгүй Ажиллагаа",
            "Үйл Ажиллагааны Гол Эрсдэлүүд",
            "Удирдлагын Арга Хэмжээний Бүртгэл",
            "Удирдлагын Зөвлөмж",
        )
        all_text = "\n".join(self._slide_text(slide) for slide in presentation.slides)
        for index, title in enumerate(expected_titles):
            self.assertIn(title, self._slide_text(presentation.slides[index]))
        self.assertIn("Ачит-Ихт Зэсийн Катодын Үйлдвэр", all_text)
        self.assertIn("Өндөр", all_text)
        self.assertIn("Дунд", all_text)
        self.assertIn("Хэрэгжиж Байна", all_text)
        self.assertIn("Дууссан", all_text)
        self.assertIn("Тодорхойлоогүй", all_text)
        self.assertNotIn("In_Progress", all_text)
        self.assertTrue(all(character in all_text for character in "ӨөҮү"))
        production_chart = next(
            shape.chart
            for shape in presentation.slides[2].shapes
            if getattr(shape, "has_chart", False)
        )
        self.assertEqual(
            [series.name for series in production_chart.series],
            ["Катодын Төлөвлөгөө", "Катодын Бодит Гүйцэтгэл"],
        )
        forbidden_product_text = (
            "Executive KPI Summary",
            "Cathode Production Trend",
            "Plant and Safety Overview",
            "Key Operational Risks",
            "Management Action Register",
            "Executive Recommendations",
            "Priority",
            "Timing",
            "In Progress",
            "Not Started",
            "Overdue",
            "Unassigned",
            "Not set",
        )
        self.assertTrue(all(token not in all_text for token in forbidden_product_text))

    def test_powerpoint_en_mn_chart_and_action_semantic_parity(self):
        english = self._generate("en")
        mongolian = self._generate("mn")
        self.assertEqual(self._chart_signature(english), self._chart_signature(mongolian))
        self.assertEqual(
            [len(slide.shapes) for slide in english.slides],
            [len(slide.shapes) for slide in mongolian.slides],
        )
        english_actions = self._slide_text(english.slides[5])
        mongolian_actions = self._slide_text(mongolian.slides[5])
        for source_text in (
            "Review the lowest-performing operational KPI",
            "Investigate production loss versus plan",
            "Үйл ажиллагаа",
            "Processing Operations",
        ):
            self.assertIn(source_text, english_actions)
        for localized_or_preserved_text in (
            "Хамгийн сул гүйцэтгэлтэй үйл ажиллагааны KPI-г хянан үзэх",
            "Үйлдвэрлэлийн төлөвлөгөөний зөрүүний шалтгааныг судлах",
            "Үйл ажиллагаа",
            "Боловсруулах үйлдвэрийн үйл ажиллагаа",
        ):
            self.assertIn(localized_or_preserved_text, mongolian_actions)
        self.assertNotIn("Review the lowest-performing operational KPI", mongolian_actions)
        self.assertNotIn("Investigate production loss versus plan", mongolian_actions)
        self.assertNotIn("Processing Operations", mongolian_actions)
        self.assertLess(
            mongolian_actions.index(
                "Хамгийн сул гүйцэтгэлтэй үйл ажиллагааны KPI-г хянан үзэх"
            ),
            mongolian_actions.index(
                "Үйлдвэрлэлийн төлөвлөгөөний зөрүүний шалтгааныг судлах"
            ),
        )

    def test_powerpoint_action_display_mappings_are_curated_and_language_scoped(self):
        action_title_mappings = {
            "Review the lowest-performing operational KPI": (
                "Хамгийн сул гүйцэтгэлтэй үйл ажиллагааны KPI-г хянан үзэх"
            ),
            "Investigate production loss versus plan": (
                "Үйлдвэрлэлийн төлөвлөгөөний зөрүүний шалтгааныг судлах"
            ),
        }
        owner_mappings = {
            "Operations": "Үйл ажиллагаа",
            "Processing Operations": "Боловсруулах үйлдвэрийн үйл ажиллагаа",
            "Үйл ажиллагаа": "Үйл ажиллагаа",
        }

        for source, expected in action_title_mappings.items():
            self.assertEqual(
                powerpoint_service._localized_action_title(source, "mn"),
                expected,
            )
            self.assertEqual(
                powerpoint_service._localized_action_title(source, "en"),
                source,
            )

        for source, expected in owner_mappings.items():
            self.assertEqual(
                powerpoint_service._localized_action_owner(source, "mn"),
                expected,
            )
            self.assertEqual(
                powerpoint_service._localized_action_owner(source, "en"),
                source,
            )

        custom_title = "Confirm customer-specific operating constraint"
        personal_owner = "B. Batbold"
        self.assertEqual(
            powerpoint_service._localized_action_title(custom_title, "mn"),
            custom_title,
        )
        self.assertEqual(
            powerpoint_service._localized_action_owner(personal_owner, "mn"),
            personal_owner,
        )


class ReportLanguagePlumbingTests(unittest.TestCase):
    tenant = {
        "company_id": 2,
        "mine_id": 2,
        "company_name": "Test Company",
        "company_name_en": "Test Company",
        "company_name_mn": "Туршилтын Компани",
        "mine_name": "Test Mine",
        "mine_name_en": "Test Mine",
        "mine_name_mn": "Туршилтын Уурхай",
        "operation_profile": "sxew_copper",
    }
    user = SimpleNamespace(id=7, full_name="Test User", email="test@example.com")

    @staticmethod
    def _capture_response(**kwargs):
        kwargs["generator"]()
        return kwargs

    def _exercise_endpoint(self, endpoint_name: str, language: str | None):
        endpoint = getattr(reports, endpoint_name)
        common_patches = [
            patch.object(reports, "_resolve_tenant", return_value=self.tenant),
            patch.object(reports, "_generate_report_response", side_effect=self._capture_response),
        ]

        if endpoint_name == "download_daily_executive_pdf":
            data_patch = patch.object(
                reports,
                "get_live_kpi_summary",
                return_value={"status": "Connected to PostgreSQL"},
            )
            generator_patch = patch.object(
                reports,
                "generate_daily_executive_pdf",
                return_value=BytesIO(b"daily"),
            )
            call_kwargs = {"mine_name": None}
        elif endpoint_name == "download_weekly_operations_pdf":
            data_patch = patch.object(
                reports,
                "get_weekly_kpi_summary",
                return_value={"status": "Connected to PostgreSQL"},
            )
            generator_patch = patch.object(
                reports,
                "generate_weekly_operations_pdf",
                return_value=BytesIO(b"weekly"),
            )
            call_kwargs = {"mine_name": None}
        elif endpoint_name == "download_monthly_kpi_pdf":
            data_patch = patch.object(
                reports,
                "get_monthly_kpi_summary",
                return_value={"status": "Connected to PostgreSQL"},
            )
            generator_patch = patch.object(
                reports,
                "generate_monthly_kpi_pdf",
                return_value=BytesIO(b"monthly"),
            )
            call_kwargs = {"mine_name": None}
        elif endpoint_name == "download_executive_excel_export":
            data_patch = patch.object(reports, "get_live_kpi_summary")
            generator_patch = patch.object(
                reports,
                "generate_executive_excel_export",
                return_value=BytesIO(b"excel"),
            )
            call_kwargs = {}
        else:
            data_patch = patch.object(reports, "get_live_kpi_summary")
            generator_patch = patch.object(
                reports,
                "generate_executive_powerpoint",
                return_value=BytesIO(b"powerpoint"),
            )
            call_kwargs = {}

        with common_patches[0], common_patches[1], data_patch, generator_patch as generator:
            response = endpoint(
                lang=language,
                db=object(),
                current_user=self.user,
                **call_kwargs,
            )
            generator_call = generator.call_args

        return response, generator_call

    def test_all_endpoints_normalize_and_forward_language(self):
        endpoint_names = (
            "download_daily_executive_pdf",
            "download_weekly_operations_pdf",
            "download_monthly_kpi_pdf",
            "download_executive_excel_export",
            "download_executive_powerpoint",
        )
        for endpoint_name in endpoint_names:
            for supplied, expected in (("en", "en"), ("MN", "mn"), (None, "en")):
                with self.subTest(endpoint=endpoint_name, language=supplied):
                    response, generator_call = self._exercise_endpoint(endpoint_name, supplied)
                    self.assertEqual(generator_call.kwargs["language"], expected)
                    self.assertIn("filename*=UTF-8''", build_content_disposition(
                        response["filename"],
                        response["ascii_filename"],
                    ))

    def test_all_endpoint_query_defaults_are_english(self):
        for endpoint_name in (
            "download_daily_executive_pdf",
            "download_weekly_operations_pdf",
            "download_monthly_kpi_pdf",
            "download_executive_excel_export",
            "download_executive_powerpoint",
        ):
            endpoint = getattr(reports, endpoint_name)
            default = inspect.signature(endpoint).parameters["lang"].default
            self.assertEqual(default.default, "en")

    def test_pdf_endpoints_carry_configured_identity_to_the_shared_resolver(self):
        for endpoint_name in (
            "download_daily_executive_pdf",
            "download_weekly_operations_pdf",
            "download_monthly_kpi_pdf",
        ):
            with self.subTest(endpoint=endpoint_name):
                _, generator_call = self._exercise_endpoint(endpoint_name, "mn")
                report_data = generator_call.args[0]
                self.assertEqual(report_data["company_name_mn"], "Туршилтын Компани")
                self.assertEqual(report_data["mine_name_mn"], "Туршилтын Уурхай")


if __name__ == "__main__":
    unittest.main()
